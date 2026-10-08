from __future__ import annotations

import copy
import difflib
import io
import json
import os
import re
import uuid
from datetime import date, datetime
from zoneinfo import ZoneInfo
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st


APP_LOGIC_VERSION = "2026-09-07_operational_ui_v8"

# Public demo mode (default on): nothing is written to or read from the server's
# disk, so visitors to a shared deployment never see each other's workspace.
# Set PUBLIC_DEMO=0 to enable local autosave and named checkpoints.
PUBLIC_DEMO = os.environ.get("PUBLIC_DEMO", "1") != "0"

# ============================================================
# CORE OPERATIONAL DEFAULTS
# ============================================================

TARGET_DAYS = 55
DEFAULT_PRIORITY_FALLBACK = 32.5
def current_planning_date() -> date:
    """Return today's date in the local NHS/Manchester timezone."""
    try:
        return datetime.now(ZoneInfo("Europe/London")).date()
    except Exception:
        return date.today()


DEFAULT_PLANNING_DATE = current_planning_date()

TIERS = ["Simple", "Moderate", "Complex"]
COMPLEXITY_OPTIONS = ["Unknown", "Simple", "Moderate", "Complex"]

# The user only needs to actively choose Not started or Completed.
# A blank source/app value is treated operationally as In progress.
STATUS_INPUT_BLANK = ""
STATUS_NOT_STARTED = "Not started"
STATUS_COMPLETED = "Completed"
STATUS_IN_PROGRESS = "In progress"
TRIAL_STATUS_INPUT_OPTIONS = [STATUS_INPUT_BLANK, STATUS_NOT_STARTED, STATUS_COMPLETED]

STATUS_COLORS = {
    STATUS_IN_PROGRESS: "#F4A261",
    STATUS_NOT_STARTED: "#E76F51",
    STATUS_COMPLETED: "#2A9D8F",
}

FINAL_DECISION_OPTIONS = [
    "",
    "Proceed as recommended",
    "Proceed - manually prioritised",
    "Hold / review",
    "Escalate",
]

APP_DIR = Path(__file__).resolve().parent
SAVED_STATES_DIR = APP_DIR / "saved_states"
AUTO_SAVE_FILE = SAVED_STATES_DIR / "_autosave.json"

DEFAULT_DURATION_RANGES = {
    "CT": {
        "Simple": (1, 5),
        "Moderate": (10, 15),
        "Complex": (20, 30),
    },
    # Aseptic is not complexity-tiered in the final model.
    "Aseptic": {
        "Simple": (30, 30),
        "Moderate": (30, 30),
        "Complex": (30, 30),
    },
    "ePMA": {
        "Simple": (2, 5),
        "Moderate": (5, 10),
        "Complex": (30, 30),
    },
}

DEFAULT_ASSUMPTIONS = {
    "target_days": 55,
    "cap_ct": 3,
    "cap_aseptic": 2,
    "cap_epma": 1,
    # Scenario A default: CT retains one capacity slot until all required
    # specialist workstreams for that trial are complete.
    "ct_capacity_mode": "continuous",
    "specialist_handoff_bdays": 2,
    "horizon_bdays": 250,
    # This is now an escalation flag only; it does NOT remove a trial from scheduling.
    "max_overdue_days": 180,
    "planning_date": DEFAULT_PLANNING_DATE,
    "aseptic_default_bdays": 30,
    "solver_time_limit_seconds": 60,
    "duration_tiers": copy.deepcopy(DEFAULT_DURATION_RANGES),
    # Rows: {workstream, start_date, end_date, available_capacity, note}
    # Used for staff absence / training / temporary reduced capacity.
    "capacity_exceptions": [],
}


# ============================================================
# DURATION + STATUS HELPERS
# ============================================================

def upper_duration(team: str, tier: str, assumptions=None) -> int:
    if assumptions is None:
        assumptions = DEFAULT_ASSUMPTIONS
    effective_tier = tier if tier in TIERS else "Complex"
    _, hi = assumptions["duration_tiers"][team][effective_tier]
    return int(hi)


def aseptic_default_duration(assumptions=None) -> int:
    if assumptions is None:
        assumptions = DEFAULT_ASSUMPTIONS
    try:
        value = int(assumptions.get("aseptic_default_bdays", 30))
    except Exception:
        value = 30
    return max(1, value)


def normalise_complexity(value: Any) -> str:
    if not _has_value(value):
        return "Unknown"
    text = str(value).strip().lower()
    mapping = {
        "simple": "Simple",
        "low": "Simple",
        "moderate": "Moderate",
        "medium": "Moderate",
        "complex": "Complex",
        "high": "Complex",
        "unknown": "Unknown",
        "not known": "Unknown",
        "tbc": "Unknown",
        "n/a": "Unknown",
        "na": "Unknown",
    }
    return mapping.get(text, "Unknown")


def default_trial_durations(tier: str = "Unknown", assumptions=None):
    if assumptions is None:
        assumptions = DEFAULT_ASSUMPTIONS
    return {
        "CT": upper_duration("CT", tier, assumptions),
        "Aseptic": aseptic_default_duration(assumptions),
        "ePMA": upper_duration("ePMA", tier, assumptions),
    }


def normalise_trial_status_input(value: Any) -> str:
    """Keep only the two explicit user choices; anything else becomes blank."""
    if not _has_value(value):
        return STATUS_INPUT_BLANK
    text = str(value).strip().lower()
    if text in {"completed", "complete", "closed", "done", "archived"}:
        return STATUS_COMPLETED
    if text in {"not started", "not-started", "not yet started", "queued", "waiting"}:
        return STATUS_NOT_STARTED
    # Explicit "in progress" is intentionally stored as blank to reduce input burden.
    return STATUS_INPUT_BLANK


def effective_trial_status(value: Any) -> str:
    raw = normalise_trial_status_input(value)
    if raw == STATUS_COMPLETED:
        return STATUS_COMPLETED
    if raw == STATUS_NOT_STARTED:
        return STATUS_NOT_STARTED
    return STATUS_IN_PROGRESS


def _milestone_on_or_before(value: Any, planning_date) -> pd.Timestamp | pd.NaT:
    """Return a normalised milestone only when it was known by the planning date.

    Future milestone dates are ignored so retrospective snapshots cannot use
    information that would not yet have been available.
    """
    ts = pd.to_datetime(value, errors="coerce")
    if pd.isna(ts):
        return pd.NaT
    ts = pd.Timestamp(ts).normalize()
    cutoff = pd.Timestamp(planning_date).normalize()
    return ts if ts <= cutoff else pd.NaT


def _business_days_elapsed(start_value: Any, planning_date) -> int:
    """Approximate CT work already elapsed using the same Mon-Fri calendar as the solver.

    The count is [start, planning_date): the planning day itself remains available
    for future work. This is an automatic estimate only; specialist workstreams
    use explicit completion milestones instead of percentage-complete estimates.
    """
    start = pd.to_datetime(start_value, errors="coerce")
    end = pd.Timestamp(planning_date).normalize()
    if pd.isna(start):
        return 0
    start = pd.Timestamp(start).normalize()
    if start >= end:
        return 0
    return int(len(pd.bdate_range(start=start, end=end - pd.Timedelta(days=1))))


# Backwards compatibility for old pages / saved states.
def normalise_trial_status(value: Any) -> str:
    return effective_trial_status(value)


# ============================================================
# BASIC VALUE HELPERS
# ============================================================

def _has_value(value: Any) -> bool:
    if value is None:
        return False
    try:
        if pd.isna(value):
            return False
    except Exception:
        pass
    text = str(value).strip()
    return bool(text) and text.lower() not in {"nan", "nat"}


def _parse_yes_no(value: Any):
    if not _has_value(value):
        return None
    text = str(value).strip().lower()
    if text in {"yes", "y", "true", "1", "required", "needed"}:
        return True
    if text in {"no", "n", "false", "0", "none", "not required", "not needed", "n/a", "na"}:
        return False
    return None


def _date_or_none(value: Any):
    """Parse dashboard/workbook dates without confusing ISO and UK day-first strings.

    Streamlit date editors may round-trip dates as ``date``/``Timestamp`` objects or
    ISO strings (YYYY-MM-DD), while uploaded workbooks often contain UK-style
    DD/MM/YYYY text. ISO values must be parsed year-first; applying ``dayfirst=True``
    to an ISO value such as 2026-09-04 can otherwise be misread as 9 April 2026.
    """
    if not _has_value(value):
        return None

    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    clean = value
    if isinstance(value, str):
        clean = re.sub(r"\s+EST\s*$", "", value.strip(), flags=re.IGNORECASE)

        # Streamlit/persisted state commonly stores ISO dates. Parse those
        # explicitly before falling back to UK day-first workbook text.
        if re.match(r"^\d{4}-\d{2}-\d{2}(?:[ T].*)?$", clean):
            parsed = pd.to_datetime(clean, errors="coerce", yearfirst=True)
        else:
            parsed = pd.to_datetime(clean, errors="coerce", dayfirst=True)
    else:
        parsed = pd.to_datetime(clean, errors="coerce")

    if pd.isna(parsed):
        return None
    return parsed.date()


def format_date_for_display(value: Any) -> str:
    """Return a UK-style DD/MM/YYYY string for a date-like value."""
    parsed = _date_or_none(value)
    return parsed.strftime("%d/%m/%Y") if parsed is not None else ""


def dates_as_date_objects(df: pd.DataFrame, columns) -> pd.DataFrame:
    """Convert selected columns to Python ``date`` objects for Streamlit/Excel display.

    Internally the optimiser can continue to use pandas Timestamps; this helper is
    presentation-only and prevents tables from showing ``00:00:00``.
    """
    out = df.copy()
    for col in columns:
        if col not in out.columns:
            continue
        parsed = pd.to_datetime(out[col], errors="coerce")
        out[col] = parsed.map(lambda x: x.date() if pd.notna(x) else None)
    return out


def sync_schedule_metadata(schedule: pd.DataFrame, trial_table: pd.DataFrame | None) -> pd.DataFrame:
    """Refresh editable metadata in an existing schedule from the live trial table.

    Notes and operational decisions do not require a mathematical re-solve.  The
    schedule view therefore hydrates them from the current editable table on every
    rerun, so a Pharmacy Note edited in another tab appears immediately.
    Complexity/duration fields are also refreshed for transparent audit of what
    numerical inputs were used at the latest solve.
    """
    if schedule is None:
        return schedule
    out = schedule.copy()
    if trial_table is None or len(trial_table) == 0 or "trial_id" not in out.columns:
        return out

    meta_cols = [
        "trial_id", "trial_status", "final_decision", "pharmacy_note", "manual_rank",
        "ct_complexity", "epma_complexity", "ct_duration_bdays",
        "aseptic_duration_bdays", "epma_duration_bdays",
        "pharmacy_setup_start_date", "aseptic_notify_date", "aseptic_bmr_authorised_date",
        "iqemo_setup_date", "iqemo_authorised_date", "greenlight_date",
    ]
    meta_cols = [c for c in meta_cols if c in trial_table.columns]
    if len(meta_cols) <= 1:
        return out

    meta = trial_table[meta_cols].copy()
    meta["trial_id"] = meta["trial_id"].astype(str)
    meta = meta.drop_duplicates("trial_id", keep="last").set_index("trial_id")
    keys = out["trial_id"].astype(str)
    for col in meta.columns:
        mapped = keys.map(meta[col])
        # Live editable metadata should replace stale copies even when blank.
        out[col] = mapped.values
    if "trial_status" in out.columns:
        out["effective_trial_status"] = out["trial_status"].apply(effective_trial_status)
    return out


def render_page_navigation(previous_page=None, next_page=None, previous_label="Previous", next_label="Next"):
    """Render simple bottom-of-page navigation for the operational workflow."""
    st.divider()
    left, right = st.columns(2)
    with left:
        if previous_page:
            if st.button(f"← {previous_label}", use_container_width=True, key=f"nav_prev_{previous_page}"):
                st.switch_page(previous_page)
    with right:
        if next_page:
            if st.button(f"{next_label} →", type="primary", use_container_width=True, key=f"nav_next_{next_page}"):
                st.switch_page(next_page)


def _number_or_none(value: Any):
    if not _has_value(value):
        return None
    x = pd.to_numeric(value, errors="coerce")
    if pd.isna(x):
        return None
    return float(x)


def _text_or_blank(value: Any) -> str:
    return str(value).strip() if _has_value(value) else ""


# ============================================================
# FLEXIBLE COLUMN MAPPING
# ============================================================

# Internal canonical field -> known exact aliases.
# Column order is irrelevant. The loader normalises punctuation, line breaks,
# spacing and capitalisation before matching.
TRACKER_CANDIDATES = {
    "row_number": ["#", "No", "No.", "Number", "Row"],
    "trial_id": [
        "Trial",
        "Trial name",
        "Trial ID",
        "Study",
        "Study name",
        "Project Short title",
        "Project short title",
    ],
    "trial_status": [
        "Trial status",
        "Status",
        "Pharmacy status",
    ],
    "actual_completion": [
        "Actual completion date",
        "Completion date",
        "Pharmacy completion date",
    ],
    "hra_approval": [
        "HRA approval",
        "HRA Approval Date",
        "HRA approval date",
        "Date HRA approved",
    ],
    "site_selected": [
        "Site selected",
        "Site selection date",
        "Date site selected",
        "Project site date site selected",
    ],
    # Audit/context only, never a substitute for Site selected.
    "site_confirmed": [
        "Date site confirmed (EDGE) MUST HAVE",
        "Date site confirmed (EDGE)",
        "Date site confirmed",
        "EDGE site confirmed",
    ],
    "pharmacy_notified": [
        "Pharmacy notified",
        "Pharmacy notification date",
        "Pharmacy notified date",
        "Pharmacy notified by RSG- start review",
        "Date pharmacy told to start",
    ],
    "pharmacy_setup_start": [
        "Pharmacy set-up start (great to have)",
        "Pharmacy set-up start",
        "Pharmacy setup start",
    ],
    "aseptic_notify": [
        "Pharmacy CT notify Aseptic Services",
        "Pharmacy CT notify Aseptic Service",
        "Pharmacy CT notify Aseptic service",
        "Aseptic notified",
    ],
    "aseptic_bmr_authorised": [
        "Date Aseptic BMR Authorised",
        "Aseptic BMR authorised",
        "Aseptic BMR Authorised",
        "BMR authorised",
    ],
    "iqemo_setup": [
        "iQemo Template set up",
        "iQemo template set up",
        "iQemo Template Setup",
        "iQemo set up",
    ],
    "iqemo_authorised": [
        "iQemo template authorised",
        "iQemo Template authorised",
        "iQemo Template Authorised",
        "iQemo authorised",
    ],
    "greenlight": [
        "Pharmacy green light",
        "Pharmacy Greenlight",
        "Pharmacy Greenlight Approval",
    ],
    "aseptic_flag": [
        "Aseptic",
        "Aseptic trial",
        "Aseptic Trial",
        "Aseptic (your tracker)",
        "Aseptic required",
    ],
    "prescribing_system": [
        "Prescribing system",
        "Prescribing system (chemo ePMA / inpatient ePMA / none)",
        "ePMA prescribing system",
        "Electronic prescribing system",
    ],
    "trial_phase": [
        "Trial phase",
        "Phase",
        "Trial phase (1, 2 or 3) MUST HAVE",
    ],
    "priority_score": [
        "Stage 2 priority score",
        "Stage 2 priority score (estimates pre-filled)",
        "Priority score",
        "Priority",
    ],
    "q4b": [
        "Q4b pharmacy & aseptics burden",
        "Q4b pharmacy & aseptics burden (great to have)",
        "Q4b",
    ],
    "ct_complexity": [
        "CT Complexity",
        "CT complexity",
        "CT complexity level",
        "Trial Complexity",
        "Trial complexity",
        "Clinical Trial Complexity",
        "Clinical Trials complexity",
    ],
    "epma_complexity": [
        "ePMA Complexity",
        "ePMA complexity",
        "ePMA complexity level",
        "Electronic prescribing complexity",
    ],
    "final_decision": [
        "Final decision",
        "Pharmacy final decision",
        "Decision",
    ],
    "manual_rank": [
        "Manual rank",
        "Override rank",
        "Pharmacy rank",
    ],
    "pharmacy_note": [
        "Pharmacy note",
        "Pharmacy notes",
        "Notes",
        "Note",
        "Comments",
        "Pharmacy comments",
    ],
}

FIELD_LABELS = {
    "row_number": "Row number (optional)",
    "trial_id": "Trial name",
    "trial_status": "Trial status",
    "actual_completion": "Actual completion date",
    "hra_approval": "HRA approval",
    "site_selected": "Site selected",
    "site_confirmed": "Site confirmed (EDGE, context only)",
    "pharmacy_notified": "Pharmacy notified",
    "pharmacy_setup_start": "Pharmacy set-up start",
    "aseptic_notify": "Pharmacy CT notify Aseptic Services",
    "aseptic_bmr_authorised": "Aseptic BMR authorised",
    "iqemo_setup": "iQemo Template set up",
    "iqemo_authorised": "iQemo template authorised",
    "greenlight": "Pharmacy green light",
    "aseptic_flag": "Aseptic required",
    "prescribing_system": "Prescribing system / ePMA route",
    "trial_phase": "Trial phase",
    "priority_score": "Stage 2 priority score",
    "q4b": "Q4b burden",
    "ct_complexity": "CT Complexity",
    "epma_complexity": "ePMA Complexity",
    "final_decision": "Final decision",
    "manual_rank": "Manual rank",
    "pharmacy_note": "Pharmacy Note",
}

# Only the trial identifier is required to construct an editable table. Missing
# solver fields are loaded as blank/Unknown and surfaced for staff review.
CORE_MAPPING_FIELDS = [
    "trial_id",
    "hra_approval",
    "site_selected",
    "pharmacy_notified",
    "pharmacy_setup_start",
    "aseptic_notify",
    "aseptic_bmr_authorised",
    "iqemo_setup",
    "iqemo_authorised",
    "greenlight",
    "aseptic_flag",
    "prescribing_system",
    "ct_complexity",
    "epma_complexity",
    "priority_score",
    "trial_status",
    "final_decision",
    "pharmacy_note",
]


def _normalise_header(value: Any) -> str:
    text = str(value if value is not None else "")
    text = text.replace("\n", " ").replace("&", " and ")
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).strip().lower()
    return " ".join(text.split())


def _alias_norms(field: str):
    return {_normalise_header(x) for x in TRACKER_CANDIDATES.get(field, [])}


def resolve_tracker_columns(df: pd.DataFrame, manual_mapping: dict | None = None):
    """Resolve canonical fields using confirmed manual mapping first, then exact aliases."""
    columns = list(df.columns)
    actual = {_normalise_header(col): col for col in columns}
    resolved = {}

    manual_mapping = manual_mapping or {}
    for field, candidates in TRACKER_CANDIDATES.items():
        chosen = manual_mapping.get(field)
        if chosen in columns:
            resolved[field] = chosen
            continue

        resolved[field] = None
        for candidate in candidates:
            match = actual.get(_normalise_header(candidate))
            if match is not None:
                resolved[field] = match
                break

    return resolved


def suggest_tracker_columns(df: pd.DataFrame, resolved: dict | None = None):
    """Return fuzzy suggestions only; suggestions are never applied automatically."""
    resolved = resolved or resolve_tracker_columns(df)
    suggestions = {}
    for field, mapped in resolved.items():
        if mapped is not None:
            continue
        best_col, best_score = None, 0.0
        aliases = list(_alias_norms(field))
        for col in df.columns:
            c = _normalise_header(col)
            for a in aliases:
                score = difflib.SequenceMatcher(None, c, a).ratio()
                if score > best_score:
                    best_col, best_score = col, score
        if best_col is not None and best_score >= 0.72:
            suggestions[field] = {"column": best_col, "score": round(best_score, 3)}
    return suggestions


def mapping_summary(df: pd.DataFrame, manual_mapping: dict | None = None):
    resolved = resolve_tracker_columns(df, manual_mapping=manual_mapping)
    suggestions = suggest_tracker_columns(df, resolved)
    rows = []
    for field in CORE_MAPPING_FIELDS:
        mapped = resolved.get(field)
        if mapped:
            status = "Detected"
            suggested = ""
        elif field in suggestions:
            status = "Needs confirmation"
            suggested = str(suggestions[field]["column"])
        else:
            status = "Not supplied"
            suggested = ""
        rows.append(
            {
                "Field": FIELD_LABELS.get(field, field),
                "Detected column": mapped or "",
                "Status": status,
                "Suggested column": suggested,
            }
        )
    return pd.DataFrame(rows), resolved, suggestions


def validate_source_mapping(df: pd.DataFrame, mapping: dict | None):
    """Validate staff-confirmed source mapping before constructing trials."""
    mapping = mapping or {}
    columns = list(df.columns)
    chosen = {k: v for k, v in mapping.items() if v not in (None, "")}

    missing_columns = sorted({str(v) for v in chosen.values() if v not in columns})
    if missing_columns:
        raise ValueError(
            "Mapped source column(s) are no longer present: " + ", ".join(missing_columns)
        )

    # One source column should not silently feed two different canonical fields.
    reverse = {}
    for field, col in chosen.items():
        reverse.setdefault(col, []).append(field)
    conflicts = {col: fields for col, fields in reverse.items() if len(fields) > 1}
    if conflicts:
        details = "; ".join(
            f"{col}: " + ", ".join(FIELD_LABELS.get(f, f) for f in fields)
            for col, fields in conflicts.items()
        )
        raise ValueError(
            "A source column has been mapped to more than one app field. "
            "Please confirm the mapping: " + details
        )

    if not chosen.get("trial_id"):
        raise ValueError("Confirm which source column contains the trial name.")


def _header_score(values) -> tuple[float, int]:
    """Score a possible header row without committing any fuzzy field mapping.

    Exact aliases receive full weight. Close-looking labels can contribute a
    smaller score purely to help locate the header row; they are still surfaced
    to staff for confirmation before the data are mapped.
    """
    normalized = [_normalise_header(v) for v in values if _has_value(v)]
    if not normalized:
        return 0.0, 0

    matched_fields = set()
    score = 0.0
    weights = {
        "trial_id": 6,
        "hra_approval": 3,
        "site_selected": 3,
        "pharmacy_notified": 3,
        "pharmacy_setup_start": 1,
        "aseptic_notify": 1,
        "aseptic_bmr_authorised": 1,
        "iqemo_setup": 1,
        "iqemo_authorised": 1,
        "greenlight": 2,
        "aseptic_flag": 2,
        "prescribing_system": 2,
        "priority_score": 1,
        "ct_complexity": 1,
        "epma_complexity": 1,
        "pharmacy_note": 1,
    }

    for field, aliases in TRACKER_CANDIDATES.items():
        alias_set = {_normalise_header(a) for a in aliases}
        if any(v in alias_set for v in normalized):
            matched_fields.add(field)
            score += weights.get(field, 0.5)
            continue

        # Fuzzy similarity is used only for structural header detection. It never
        # becomes an automatic source-column mapping.
        best = 0.0
        for v in normalized:
            for alias in alias_set:
                best = max(best, difflib.SequenceMatcher(None, v, alias).ratio())
        if best >= 0.84:
            matched_fields.add(field)
            score += weights.get(field, 0.5) * 0.55

    return score, len(matched_fields)


def _read_excel_layout(file_bytes: bytes, sheet_name=None, header_row=None):
    bio = io.BytesIO(file_bytes)
    xl = pd.ExcelFile(bio)
    sheet_names = xl.sheet_names

    if sheet_name is not None:
        if sheet_name not in sheet_names:
            raise ValueError(f"Sheet '{sheet_name}' was not found.")
        selected_sheet = sheet_name
        selected_header = int(header_row or 0)
    else:
        best = None
        for sheet in sheet_names:
            try:
                preview = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, header=None, nrows=25)
            except Exception:
                continue
            for row_idx in range(len(preview)):
                score, nmatch = _header_score(preview.iloc[row_idx].tolist())
                candidate = (score, nmatch, -row_idx, sheet, row_idx)
                if best is None or candidate > best:
                    best = candidate
        if best is None or best[0] <= 0:
            # Safe fallback: first sheet, first row; user can change layout manually.
            selected_sheet, selected_header = sheet_names[0], 0
        else:
            selected_sheet, selected_header = best[3], best[4]

    df = pd.read_excel(
        io.BytesIO(file_bytes),
        sheet_name=selected_sheet,
        header=selected_header,
    )
    # Preserve named columns even when every data cell is blank. Operational
    # templates may deliberately include future-input fields such as CT/ePMA
    # Complexity with no values yet. Only truly unnamed empty spacer columns are
    # removed.
    drop_cols = [
        c for c in df.columns
        if str(c).strip().lower().startswith("unnamed:") and df[c].isna().all()
    ]
    if drop_cols:
        df = df.drop(columns=drop_cols)
    df = df.copy()
    return df, {
        "sheet_names": sheet_names,
        "sheet_name": selected_sheet,
        "header_row_zero_based": int(selected_header),
        "header_row_excel": int(selected_header) + 1,
    }


def _read_csv_layout(file_bytes: bytes, header_row=None):
    preview = pd.read_csv(io.BytesIO(file_bytes), header=None, nrows=25)
    if header_row is None:
        best = None
        for row_idx in range(len(preview)):
            score, nmatch = _header_score(preview.iloc[row_idx].tolist())
            candidate = (score, nmatch, -row_idx, row_idx)
            if best is None or candidate > best:
                best = candidate
        selected_header = best[3] if best and best[0] > 0 else 0
    else:
        selected_header = int(header_row)
    df = pd.read_csv(io.BytesIO(file_bytes), header=selected_header)
    drop_cols = [
        c for c in df.columns
        if str(c).strip().lower().startswith("unnamed:") and df[c].isna().all()
    ]
    if drop_cols:
        df = df.drop(columns=drop_cols)
    df = df.copy()
    return df, {
        "sheet_names": ["CSV"],
        "sheet_name": "CSV",
        "header_row_zero_based": int(selected_header),
        "header_row_excel": int(selected_header) + 1,
    }


def inspect_operational_file(file_obj, sheet_name=None, header_row=None):
    """
    Detect sheet/header flexibly and inspect column mapping.

    No fixed sheet name, header row, column position or portfolio size is required.
    """
    filename = getattr(file_obj, "name", "uploaded_file")
    if hasattr(file_obj, "getvalue"):
        file_bytes = file_obj.getvalue()
    else:
        file_obj.seek(0)
        file_bytes = file_obj.read()

    if filename.lower().endswith(".csv"):
        df, layout = _read_csv_layout(file_bytes, header_row=header_row)
        source_type = "CSV operational file"
    else:
        df, layout = _read_excel_layout(file_bytes, sheet_name=sheet_name, header_row=header_row)
        source_type = "Excel operational workbook"

    summary, resolved, suggestions = mapping_summary(df)
    layout.update(
        {
            "filename": filename,
            "source_type": source_type,
            "resolved": resolved,
            "suggestions": suggestions,
            "columns": [str(c) for c in df.columns],
            "mapping_summary": summary,
        }
    )
    df.attrs["source_type"] = source_type
    df.attrs["layout_info"] = layout
    return df, layout


def load_operational_tracker(file_obj, manual_mapping=None, sheet_name=None, header_row=None):
    """Backwards-compatible one-call loader using the flexible inspector."""
    df, info = inspect_operational_file(file_obj, sheet_name=sheet_name, header_row=header_row)
    if manual_mapping:
        df.attrs["column_mapping"] = dict(manual_mapping)
    else:
        df.attrs["column_mapping"] = dict(info.get("resolved") or {})
    if resolve_tracker_columns(df, df.attrs.get("column_mapping")).get("trial_id") is None:
        raise ValueError(
            "A trial-name column could not be identified confidently. Review the column mapping and confirm which column contains the trial name."
        )
    return df


# ============================================================
# ROUTING
# ============================================================

def infer_aseptic_route(row, resolved):
    col = resolved.get("aseptic_flag")
    if col:
        parsed = _parse_yes_no(row.get(col))
        if parsed is True:
            return "Yes"
        if parsed is False:
            return "No"
    return "Unknown"


def normalise_prescribing_system(value):
    if not _has_value(value):
        return ""
    text = str(value).strip()
    low = text.lower()
    if low in {"unknown", "not yet entered", "not entered", "missing"}:
        return ""
    if low in {"chemo epma", "chemo e-pma", "chemo", "iqemo", "i-qemo"}:
        return "chemo ePMA"
    if low in {"inpatient epma", "inpatient e-pma", "inpatient"}:
        return "inpatient ePMA"
    if low in {"no", "none", "n/a", "na", "not applicable", "no - epma not required", "no – epma not required"}:
        return "none"
    return text


def prescribing_system_to_epma_required(value):
    system = normalise_prescribing_system(value)
    if system in {"chemo ePMA", "inpatient ePMA"}:
        return "Yes"
    if system == "none":
        return "No"
    return "Unknown"


def infer_epma_route(row, resolved):
    col = resolved.get("prescribing_system")
    if col and _has_value(row.get(col)):
        return prescribing_system_to_epma_required(row.get(col))
    return "Unknown"


# ============================================================
# EDITABLE INTERNAL TRIAL TABLE
# ============================================================

def blank_trial_row(assumptions=None):
    assumptions = assumptions or DEFAULT_ASSUMPTIONS
    return {
        "trial_id": "",
        "trial_status": "",
        "hra_date": None,
        "site_date": None,
        "site_confirmed_date": None,
        "pharmacy_notified_date": None,
        "pharmacy_setup_start_date": None,
        "aseptic_notify_date": None,
        "aseptic_bmr_authorised_date": None,
        "iqemo_setup_date": None,
        "iqemo_authorised_date": None,
        "greenlight_date": None,
        "aseptic_required": "Unknown",
        "prescribing_system": "",
        "epma_required": "Unknown",
        "trial_phase": "",
        "priority_score": None,
        "priority_source": "Missing - cohort mean/equal fallback used at scheduling",
        "ct_complexity": "Unknown",
        "epma_complexity": "Unknown",
        "ct_duration_bdays": upper_duration("CT", "Unknown", assumptions),
        "aseptic_duration_override_bdays": None,
        "aseptic_duration_bdays": 0,
        "epma_duration_bdays": 0,
        "manual_rank": None,
        "final_decision": "",
        "pharmacy_note": "",
        "q4b_burden": None,
    }


def _optional_value(row, col):
    return row.get(col) if col else None


def build_operational_trial_table(raw_df, assumptions=None, mapping=None):
    assumptions = assumptions or DEFAULT_ASSUMPTIONS
    mapping = mapping or raw_df.attrs.get("column_mapping") or {}
    if mapping:
        validate_source_mapping(raw_df, mapping)
    resolved = resolve_tracker_columns(raw_df, manual_mapping=mapping)
    trial_col = resolved.get("trial_id")
    if trial_col is None:
        raise ValueError("Confirm which source column contains the trial name before loading.")

    working = raw_df.copy()

    # Row numbering is optional and is NEVER used as a portfolio-size rule.
    # If a numbering column exists, only obvious example/template rows are removed;
    # valid trial rows with a blank/non-numeric number are still retained.
    number_col = resolved.get("row_number")
    if number_col and number_col in working.columns:
        number_text = working[number_col].astype(str).str.strip().str.lower()
        trial_text = working[trial_col].astype(str).str.strip().str.lower()
        obvious_example = (
            number_text.isin({"e.g.", "eg", "example", "sample"})
            | trial_text.str.contains(r"\bexample\b|illustrative example", regex=True, na=False)
        )
        working = working.loc[~obvious_example].copy()

    rows = []
    for _, source_row in working.iterrows():
        trial_name = _text_or_blank(source_row.get(trial_col))
        if not trial_name:
            continue

        aseptic = infer_aseptic_route(source_row, resolved)
        prescribing_system = normalise_prescribing_system(
            _optional_value(source_row, resolved.get("prescribing_system"))
        )
        epma = prescribing_system_to_epma_required(prescribing_system)

        ct_complexity = normalise_complexity(
            _optional_value(source_row, resolved.get("ct_complexity"))
        )
        if epma == "Yes":
            epma_complexity = normalise_complexity(
                _optional_value(source_row, resolved.get("epma_complexity"))
            )
        elif epma == "No":
            epma_complexity = "N/A"
        else:
            epma_complexity = "Unknown"

        priority = _number_or_none(_optional_value(source_row, resolved.get("priority_score")))
        if priority is not None and priority <= 0:
            priority = None

        manual_rank = _number_or_none(_optional_value(source_row, resolved.get("manual_rank")))
        if manual_rank is not None and manual_rank > 0:
            manual_rank = int(round(manual_rank))
        else:
            manual_rank = None

        trial_status = normalise_trial_status_input(
            _optional_value(source_row, resolved.get("trial_status"))
        )

        ct_duration = upper_duration("CT", ct_complexity, assumptions)
        aseptic_duration = aseptic_default_duration(assumptions) if aseptic == "Yes" else 0
        epma_duration = upper_duration("ePMA", epma_complexity, assumptions) if epma == "Yes" else 0

        rows.append(
            {
                "trial_id": trial_name,
                "trial_status": trial_status,
                "hra_date": _date_or_none(_optional_value(source_row, resolved.get("hra_approval"))),
                "site_date": _date_or_none(_optional_value(source_row, resolved.get("site_selected"))),
                "site_confirmed_date": _date_or_none(_optional_value(source_row, resolved.get("site_confirmed"))),
                "pharmacy_notified_date": _date_or_none(_optional_value(source_row, resolved.get("pharmacy_notified"))),
                "pharmacy_setup_start_date": _date_or_none(_optional_value(source_row, resolved.get("pharmacy_setup_start"))),
                "aseptic_notify_date": _date_or_none(_optional_value(source_row, resolved.get("aseptic_notify"))),
                "aseptic_bmr_authorised_date": _date_or_none(_optional_value(source_row, resolved.get("aseptic_bmr_authorised"))),
                "iqemo_setup_date": _date_or_none(_optional_value(source_row, resolved.get("iqemo_setup"))),
                "iqemo_authorised_date": _date_or_none(_optional_value(source_row, resolved.get("iqemo_authorised"))),
                "greenlight_date": _date_or_none(_optional_value(source_row, resolved.get("greenlight"))),
                "aseptic_required": aseptic,
                "prescribing_system": prescribing_system,
                "epma_required": epma,
                "trial_phase": _text_or_blank(_optional_value(source_row, resolved.get("trial_phase"))),
                "priority_score": priority,
                "priority_source": "Uploaded score" if priority is not None else "Missing - cohort mean/equal fallback used at scheduling",
                "ct_complexity": ct_complexity,
                "epma_complexity": epma_complexity,
                "ct_duration_bdays": ct_duration,
                "aseptic_duration_override_bdays": None,
                "aseptic_duration_bdays": aseptic_duration,
                "epma_duration_bdays": epma_duration,
                "manual_rank": manual_rank,
                "final_decision": _text_or_blank(_optional_value(source_row, resolved.get("final_decision"))),
                "pharmacy_note": _text_or_blank(_optional_value(source_row, resolved.get("pharmacy_note"))),
                "q4b_burden": _number_or_none(_optional_value(source_row, resolved.get("q4b"))),
            }
        )

    out = pd.DataFrame(rows)
    if out.empty:
        return out
    # Duplicate trial names are operationally dangerous because edits/overrides
    # are keyed by trial name. Surface them explicitly rather than silently merge.
    if out["trial_id"].astype(str).str.strip().duplicated().any():
        dupes = out.loc[out["trial_id"].astype(str).str.strip().duplicated(keep=False), "trial_id"].tolist()
        raise ValueError("Duplicate trial names were found: " + ", ".join(map(str, dupes)))
    return normalise_edited_trials(None, out, assumptions)


def normalise_edited_trials(old_df, edited_df, assumptions=None):
    assumptions = assumptions or DEFAULT_ASSUMPTIONS
    edited = edited_df.copy().reset_index(drop=True)
    old = old_df.copy().reset_index(drop=True) if old_df is not None else pd.DataFrame()
    if edited.empty:
        return edited

    required_defaults = blank_trial_row(assumptions)
    for col, default in required_defaults.items():
        if col not in edited.columns:
            edited[col] = default

    for idx in edited.index:
        edited.at[idx, "trial_id"] = str(edited.at[idx, "trial_id"]).strip()
        edited.at[idx, "trial_status"] = normalise_trial_status_input(edited.at[idx, "trial_status"])

        # Dates
        for col in [
            "hra_date", "site_date", "site_confirmed_date", "pharmacy_notified_date",
            "pharmacy_setup_start_date", "aseptic_notify_date", "aseptic_bmr_authorised_date",
            "iqemo_setup_date", "iqemo_authorised_date", "greenlight_date",
        ]:
            edited.at[idx, col] = _date_or_none(edited.at[idx, col])

        # Priority
        p = _number_or_none(edited.at[idx, "priority_score"])
        if p is None or p <= 0:
            edited.at[idx, "priority_score"] = None
            edited.at[idx, "priority_source"] = "Missing - cohort mean/equal fallback used at scheduling"
        else:
            old_p = _number_or_none(old.at[idx, "priority_score"]) if idx in old.index and "priority_score" in old.columns else None
            edited.at[idx, "priority_score"] = float(p)
            if old_p is None or float(old_p) != float(p):
                edited.at[idx, "priority_source"] = "Manual entry in app"

        # CT complexity always determines the numerical upper-bound duration.
        ct = normalise_complexity(edited.at[idx, "ct_complexity"])
        edited.at[idx, "ct_complexity"] = ct
        edited.at[idx, "ct_duration_bdays"] = upper_duration("CT", ct, assumptions)

        # Aseptic route + duration.
        a = str(edited.at[idx, "aseptic_required"]).strip().title()
        if a not in {"Yes", "No", "Unknown"}:
            a = "Unknown"
        edited.at[idx, "aseptic_required"] = a
        if a == "Yes":
            override = _number_or_none(edited.at[idx, "aseptic_duration_override_bdays"])
            if override is not None and override > 0:
                override = int(round(override))
                edited.at[idx, "aseptic_duration_override_bdays"] = override
                edited.at[idx, "aseptic_duration_bdays"] = override
            else:
                edited.at[idx, "aseptic_duration_override_bdays"] = None
                edited.at[idx, "aseptic_duration_bdays"] = aseptic_default_duration(assumptions)
        else:
            edited.at[idx, "aseptic_duration_override_bdays"] = None
            edited.at[idx, "aseptic_duration_bdays"] = 0

        # ePMA route comes from prescribing system only.
        system = normalise_prescribing_system(edited.at[idx, "prescribing_system"])
        edited.at[idx, "prescribing_system"] = system
        e_required = prescribing_system_to_epma_required(system)
        edited.at[idx, "epma_required"] = e_required
        if e_required == "Yes":
            e = normalise_complexity(edited.at[idx, "epma_complexity"])
            edited.at[idx, "epma_complexity"] = e
            edited.at[idx, "epma_duration_bdays"] = upper_duration("ePMA", e, assumptions)
        elif e_required == "No":
            edited.at[idx, "epma_complexity"] = "N/A"
            edited.at[idx, "epma_duration_bdays"] = 0
        else:
            edited.at[idx, "epma_complexity"] = "Unknown"
            edited.at[idx, "epma_duration_bdays"] = 0

        rank = _number_or_none(edited.at[idx, "manual_rank"])
        edited.at[idx, "manual_rank"] = int(round(rank)) if rank is not None and rank > 0 else None
        edited.at[idx, "final_decision"] = _text_or_blank(edited.at[idx, "final_decision"])
        edited.at[idx, "pharmacy_note"] = _text_or_blank(edited.at[idx, "pharmacy_note"])

    return edited


def scheduling_queue_ids(trial_table, assumptions=None):
    """Return trial IDs currently eligible for the scheduling optimiser."""
    operational = build_operational_view(trial_table, assumptions or DEFAULT_ASSUMPTIONS)
    if operational is None or len(operational) == 0:
        return []
    ready = operational.loc[
        operational.get("model_ready", pd.Series(False, index=operational.index)).fillna(False).astype(bool)
    ].copy()
    return ready["trial_id"].astype(str).tolist()


def manual_queue_order(trial_table, assumptions=None):
    """Return a complete saved manual queue order, otherwise an empty list.

    Manual override is intentionally treated as a queue move rather than a sparse
    label. A complete 1..N order avoids the ambiguous case where a lone rank of 5
    would otherwise behave exactly like rank 1.
    """
    queue_ids = scheduling_queue_ids(trial_table, assumptions)
    if not queue_ids or trial_table is None or "manual_rank" not in trial_table.columns:
        return []
    table = trial_table.copy()
    table["_trial_id_str"] = table["trial_id"].astype(str)
    q = table.loc[table["_trial_id_str"].isin(queue_ids), ["_trial_id_str", "manual_rank"]].copy()
    q["manual_rank"] = pd.to_numeric(q["manual_rank"], errors="coerce")
    if len(q) != len(queue_ids) or q["manual_rank"].isna().any():
        return []
    q["manual_rank"] = q["manual_rank"].round().astype(int)
    if sorted(q["manual_rank"].tolist()) != list(range(1, len(q) + 1)):
        return []
    return q.sort_values("manual_rank")["_trial_id_str"].tolist()


def apply_manual_queue_move(trial_table, trial_id, target_rank, assumptions=None, baseline_order=None):
    """Move one schedulable trial to a queue position and renumber the queue.

    The first manual move creates a complete 1..N queue order. Existing complete
    manual order is used as the starting point; otherwise ``baseline_order`` (for
    example the latest optimiser recommendation) is used. Remaining IDs are
    appended deterministically if the supplied baseline is incomplete.
    """
    if trial_table is None or len(trial_table) == 0:
        raise ValueError("No trial table is available.")

    assumptions = assumptions or DEFAULT_ASSUMPTIONS
    queue_ids = scheduling_queue_ids(trial_table, assumptions)
    target = str(trial_id)
    if target not in queue_ids:
        raise ValueError("Manual queue override can only be applied to trials currently in the scheduling queue.")

    try:
        rank = int(target_rank)
    except Exception as exc:
        raise ValueError("Manual queue position must be a positive whole number.") from exc
    if rank < 1 or rank > len(queue_ids):
        raise ValueError(f"Manual queue position must be between 1 and {len(queue_ids)}.")

    current_order = manual_queue_order(trial_table, assumptions)
    if current_order:
        order = current_order
    else:
        supplied = [str(x) for x in (baseline_order or [])]
        order = []
        for tid in supplied + queue_ids:
            if tid in queue_ids and tid not in order:
                order.append(tid)

    order = [tid for tid in order if tid != target]
    order.insert(rank - 1, target)

    updated = trial_table.copy().reset_index(drop=True)
    updated["_trial_id_str"] = updated["trial_id"].astype(str)
    # A manual override describes the *current* schedulable queue. Clear stale
    # ranks first so future/not-ready trials cannot collide with the live order.
    updated["manual_rank"] = None
    rank_map = {tid: pos for pos, tid in enumerate(order, start=1)}
    for idx in updated.index:
        tid = updated.at[idx, "_trial_id_str"]
        if tid in rank_map:
            updated.at[idx, "manual_rank"] = int(rank_map[tid])
    updated = updated.drop(columns=["_trial_id_str"])
    return updated


def clear_manual_queue_order(trial_table):
    """Remove manual queue ordering and return scheduling control to the optimiser."""
    if trial_table is None:
        return trial_table
    updated = trial_table.copy().reset_index(drop=True)
    if "manual_rank" in updated.columns:
        updated["manual_rank"] = None
    return updated


def validate_manual_ranks(trial_table, assumptions=None):
    """Validate positive override ranks in the current scheduling queue."""
    if trial_table is None or len(trial_table) == 0 or "manual_rank" not in trial_table.columns:
        return

    if assumptions is not None:
        queue_ids = set(scheduling_queue_ids(trial_table, assumptions))
        active = trial_table.loc[trial_table["trial_id"].astype(str).isin(queue_ids)].copy()
    else:
        active = trial_table.copy()
        if "trial_status" in active.columns:
            completed_mask = active["trial_status"].apply(
                lambda value: effective_trial_status(normalise_trial_status_input(value)) == STATUS_COMPLETED
            )
            active = active.loc[~completed_mask].copy()

    ranks = pd.to_numeric(active["manual_rank"], errors="coerce")
    ranks = ranks[ranks > 0].round().astype(int)
    duplicates = sorted(ranks[ranks.duplicated(keep=False)].unique().tolist())
    if duplicates:
        raise ValueError(
            "Manual queue positions must be unique. Duplicate position(s): "
            + ", ".join(map(str, duplicates))
            + ". Use the move control to reposition a trial; the app will renumber the rest automatically."
        )


# ============================================================
# REAL-DATE OPERATIONAL VIEW
# ============================================================

def build_operational_view(trial_table, assumptions=None):
    """Build the planning-date portfolio with milestone-aware remaining work.

    Historical milestones are treated as fixed facts only when dated on/before the
    selected planning date. Completed Aseptic/ePMA workstreams are removed from
    future resource demand; incomplete specialist workstreams retain their full
    selected duration as a conservative residual assumption. CT is deliberately
    more conservative: Pharmacy set-up start proves that CT work has started, but
    the tracker has no dedicated CT-completion milestone. Elapsed time is therefore
    NOT treated as completed CT work. Until the whole trial is completed/archived,
    the selected CT duration is retained as a planning allowance.
    """
    assumptions = assumptions or DEFAULT_ASSUMPTIONS
    if trial_table is None or len(trial_table) == 0:
        return pd.DataFrame()

    target_days = int(assumptions.get("target_days", 55))
    escalation_days = int(assumptions.get("max_overdue_days", 180))
    planning_date = pd.Timestamp(assumptions.get("planning_date", current_planning_date())).normalize()

    priority_values = pd.to_numeric(trial_table.get("priority_score", pd.Series(dtype=float)), errors="coerce")
    positive = priority_values[priority_values > 0]
    priority_mean = float(positive.mean()) if len(positive) else float(DEFAULT_PRIORITY_FALLBACK)
    no_observed_priority = len(positive) == 0

    rows = []
    for _, row in trial_table.iterrows():
        hra = pd.to_datetime(row.get("hra_date"), errors="coerce")
        site = pd.to_datetime(row.get("site_date"), errors="coerce")
        notified = pd.to_datetime(row.get("pharmacy_notified_date"), errors="coerce")

        # Progress milestones: future dates are deliberately ignored.
        setup_start = _milestone_on_or_before(row.get("pharmacy_setup_start_date"), planning_date)
        aseptic_notify = _milestone_on_or_before(row.get("aseptic_notify_date"), planning_date)
        aseptic_bmr = _milestone_on_or_before(row.get("aseptic_bmr_authorised_date"), planning_date)
        iqemo_setup = _milestone_on_or_before(row.get("iqemo_setup_date"), planning_date)
        iqemo_authorised = _milestone_on_or_before(row.get("iqemo_authorised_date"), planning_date)
        greenlight = _milestone_on_or_before(row.get("greenlight_date"), planning_date)

        raw_status = normalise_trial_status_input(row.get("trial_status"))
        explicit_completed = effective_trial_status(raw_status) == STATUS_COMPLETED
        completed_from_greenlight = pd.notna(greenlight)
        completed = bool(explicit_completed or completed_from_greenlight)

        priority = pd.to_numeric(row.get("priority_score"), errors="coerce")
        if pd.notna(priority) and float(priority) > 0:
            effective_priority = float(priority)
            priority_imputed = False
        else:
            effective_priority = priority_mean
            priority_imputed = True

        if pd.notna(hra) and pd.notna(site):
            clock_start = max(hra, site)
            due_date = clock_start + pd.Timedelta(days=target_days)
        else:
            clock_start = pd.NaT
            due_date = pd.NaT

        stage2_started = bool(pd.notna(clock_start) and clock_start <= planning_date)
        not_yet_stage2 = bool(pd.notna(clock_start) and clock_start > planning_date)

        # A blank notification after Stage 2 starts means the trial is not yet
        # available to Pharmacy (or the notification has not yet been recorded).
        awaiting_pharmacy = bool(
            stage2_started
            and (pd.isna(notified) or pd.Timestamp(notified).normalize() > planning_date)
        )
        pharmacy_ready = bool(
            stage2_started
            and pd.notna(notified)
            and pd.Timestamp(notified).normalize() <= planning_date
        )

        if pd.notna(clock_start) and pd.notna(notified):
            scheduling_ready = max(clock_start, notified)
            operational_release = max(scheduling_ready, planning_date)
        else:
            scheduling_ready = pd.NaT
            operational_release = pd.NaT

        days_to_due = int((due_date - planning_date).days) if pd.notna(due_date) else None
        overdue_days = max(0, -days_to_due) if days_to_due is not None else None
        escalation_needed = bool(overdue_days is not None and overdue_days > escalation_days)

        aseptic_required = row.get("aseptic_required") == "Yes"
        epma_required = row.get("epma_required") == "Yes"

        # Completion milestones are authoritative for specialist branches.
        aseptic_completed = bool(aseptic_required and (completed or pd.notna(aseptic_bmr)))
        epma_completed = bool(epma_required and (completed or pd.notna(iqemo_authorised)))

        aseptic_started = bool(
            aseptic_required
            and (pd.notna(aseptic_notify) or pd.notna(aseptic_bmr))
        )
        epma_started = bool(
            epma_required
            and (pd.notna(iqemo_setup) or pd.notna(iqemo_authorised))
        )
        ct_started = bool(pd.notna(setup_start))

        ct_full = pd.to_numeric(row.get("ct_duration_bdays"), errors="coerce")
        ct_full = int(ct_full) if pd.notna(ct_full) and float(ct_full) >= 0 else 0
        aseptic_full = pd.to_numeric(row.get("aseptic_duration_bdays"), errors="coerce")
        aseptic_full = int(aseptic_full) if pd.notna(aseptic_full) and float(aseptic_full) >= 0 else 0
        epma_full = pd.to_numeric(row.get("epma_duration_bdays"), errors="coerce")
        epma_full = int(epma_full) if pd.notna(epma_full) and float(epma_full) >= 0 else 0

        # CT has no dedicated completion milestone in the source tracker.
        # A recorded Pharmacy set-up start is evidence that CT work has started,
        # but elapsed calendar/business time is not evidence that the selected
        # CT workload has been completed.  To avoid false precision, started CT
        # work retains the FULL selected duration as a conservative planning
        # allowance until the whole trial is completed/archived.
        ct_elapsed_context = (
            _business_days_elapsed(setup_start, planning_date)
            if ct_started else 0
        )
        if completed:
            ct_remaining = 0
        else:
            ct_remaining = ct_full

        if not aseptic_required or aseptic_completed:
            aseptic_remaining = 0
        else:
            # Conservative rule: progress-start evidence alone does not imply a
            # percentage complete. Full residual is retained until BMR authorised.
            aseptic_remaining = aseptic_full

        if not epma_required or epma_completed:
            epma_remaining = 0
        else:
            # Same conservative rule: full residual remains until iQemo authorised.
            epma_remaining = epma_full

        if completed:
            ct_progress = "Completed"
        elif ct_started:
            ct_progress = "Started / completion not recorded"
        else:
            ct_progress = "Not started"

        if not aseptic_required:
            aseptic_progress = "N/A"
        elif aseptic_completed:
            aseptic_progress = "Completed"
        elif aseptic_started:
            aseptic_progress = "In progress"
        else:
            aseptic_progress = "Not started"

        if not epma_required:
            epma_progress = "N/A"
        elif epma_completed:
            epma_progress = "Completed"
        elif epma_started:
            epma_progress = "In progress"
        else:
            epma_progress = "Not started"

        progress_seen = bool(
            ct_started or aseptic_started or epma_started or aseptic_completed or epma_completed
        )
        if completed:
            effective_status = STATUS_COMPLETED
        elif progress_seen:
            effective_status = STATUS_IN_PROGRESS
        elif pharmacy_ready:
            effective_status = STATUS_NOT_STARTED
        else:
            # Outside the live queue, preserve a meaningful lightweight label.
            effective_status = (
                STATUS_NOT_STARTED
                if effective_trial_status(raw_status) == STATUS_NOT_STARTED
                else STATUS_IN_PROGRESS
            )

        missing = []
        if not completed:
            if pd.isna(hra):
                missing.append("HRA approval")
            if pd.isna(site):
                missing.append("Site selected")
            if row.get("aseptic_required") == "Unknown":
                missing.append("Aseptic route")
            if row.get("epma_required") == "Unknown":
                missing.append("Prescribing system / ePMA route")

        remaining_parts = []
        if ct_remaining > 0:
            remaining_parts.append(f"CT {ct_remaining} working day(s)")
        if aseptic_remaining > 0:
            remaining_parts.append(f"Aseptic {aseptic_remaining} working day(s)")
        if epma_remaining > 0:
            remaining_parts.append(f"ePMA {epma_remaining} working day(s)")
        has_remaining_modelled_work = bool(remaining_parts)

        no_modelled_work_remaining = bool(
            not completed
            and pharmacy_ready
            and len(missing) == 0
            and not has_remaining_modelled_work
        )

        model_ready = bool(
            not completed
            and len(missing) == 0
            and pharmacy_ready
            and has_remaining_modelled_work
        )
        needs_data_attention = bool(
            not completed
            and not not_yet_stage2
            and not awaiting_pharmacy
            and len(missing) > 0
        )

        reason_parts = []
        if completed:
            if completed_from_greenlight:
                reason_parts.append(
                    f"Pharmacy Green Light recorded on {pd.Timestamp(greenlight).strftime('%d %b %Y')}; "
                    "the trial is completed and excluded from future scheduling."
                )
            else:
                reason_parts.append("Marked Completed; the trial is archived and excluded from future scheduling.")
        elif pd.isna(hra) or pd.isna(site):
            if pd.isna(hra):
                reason_parts.append("HRA approval is missing, so the 55-day clock cannot be determined.")
            if pd.isna(site):
                reason_parts.append("Site selected is missing, so the 55-day clock cannot be determined.")
        elif not_yet_stage2:
            reason_parts.append(
                f"Stage 2 has not started by the selected planning date; clock start is {pd.Timestamp(clock_start).strftime('%d %b %Y')}."
            )
        elif awaiting_pharmacy:
            if pd.isna(notified):
                reason_parts.append(
                    "Stage 2 is active but Pharmacy notification has not yet been recorded; "
                    "the trial is excluded until Pharmacy receives/records the notification."
                )
            else:
                reason_parts.append(
                    f"Stage 2 is active, but Pharmacy notification is recorded for {pd.Timestamp(notified).strftime('%d %b %Y')}."
                )

        if not completed:
            if row.get("aseptic_required") == "Unknown":
                reason_parts.append("Aseptic routing needs confirmation.")
            if row.get("epma_required") == "Unknown":
                reason_parts.append("Prescribing system / ePMA routing needs confirmation.")

        if no_modelled_work_remaining:
            reason_parts.append(
                "All modelled CT/Aseptic/ePMA work is complete or estimated complete, "
                "but Pharmacy Green Light is not recorded. Excluded from the optimiser for final operational review."
            )
        elif model_ready:
            reason_parts.append(
                "Included in the optimiser. Remaining modelled work: " + ", ".join(remaining_parts) + "."
            )

        if escalation_needed and not completed:
            reason_parts.append(
                f"Escalation flag: {int(overdue_days)} calendar days overdue at the planning date. "
                + ("The trial remains in scheduling." if model_ready else "")
            )

        if completed:
            snapshot_category = "Completed / archived"
            data_status = "Completed"
        elif not_yet_stage2:
            snapshot_category = "Not yet Stage 2 started"
            data_status = "Not yet Stage 2"
        elif awaiting_pharmacy:
            snapshot_category = "Awaiting Pharmacy notification"
            data_status = "Awaiting Pharmacy"
        elif needs_data_attention:
            snapshot_category = "Needs data attention"
            data_status = "Needs input"
        elif no_modelled_work_remaining:
            snapshot_category = "Final completion review"
            data_status = "No modelled work remaining"
        elif model_ready:
            snapshot_category = "Snapshot queue"
            if escalation_needed:
                data_status = "Overdue - escalation recommended"
            elif days_to_due is not None and days_to_due < 0:
                data_status = "Overdue"
            elif days_to_due is not None and days_to_due <= 14:
                data_status = "Attention"
            else:
                data_status = "Ready"
        else:
            snapshot_category = "Needs data attention"
            data_status = "Needs review"

        queue_type = (
            "Existing / in progress"
            if model_ready and effective_status == STATUS_IN_PROGRESS
            else ("New / not started" if model_ready else "")
        )

        if model_ready:
            scheduling_action = (
                f"Included — optimise remaining work ({', '.join(remaining_parts)})."
            )
        elif completed:
            scheduling_action = "Excluded — completed / archived."
        elif no_modelled_work_remaining:
            scheduling_action = "Excluded — no modelled work remains; review final Green Light."
        elif awaiting_pharmacy:
            scheduling_action = "Excluded — awaiting Pharmacy notification."
        elif not_yet_stage2:
            scheduling_action = "Excluded — Stage 2 not yet started."
        elif needs_data_attention:
            scheduling_action = "Excluded — required scheduling input is missing."
        else:
            scheduling_action = "Excluded — operational review required."

        output = row.to_dict()
        output.update(
            {
                "trial_status": raw_status,
                "effective_trial_status": effective_status,
                "clock_start": clock_start,
                "due_date": due_date,
                "scheduling_ready_date": scheduling_ready,
                "operational_release_date": operational_release,
                "days_to_55_day_due": days_to_due,
                "overdue_days": overdue_days,
                "escalation_needed": escalation_needed,
                "stage2_started_by_planning_date": stage2_started,
                "not_yet_stage2_started": not_yet_stage2,
                "awaiting_pharmacy_notification": awaiting_pharmacy,
                "ready_by_planning_date": pharmacy_ready,
                "effective_priority_score": effective_priority,
                "priority_is_mean_imputed": priority_imputed,
                "priority_fallback_is_equal": no_observed_priority,
                "model_ready": model_ready,
                "needs_data_attention": needs_data_attention,
                "needs_operational_review": bool((escalation_needed and model_ready) or no_modelled_work_remaining),
                "snapshot_category": snapshot_category,
                "missing_required_input": "; ".join(missing),
                "scheduling_block_reason": " ".join(reason_parts),
                "scheduling_action": scheduling_action,
                "data_status": data_status,
                "queue_type": queue_type,
                "completed_from_greenlight": completed_from_greenlight,
                "no_modelled_work_remaining": no_modelled_work_remaining,
                "has_remaining_modelled_work": has_remaining_modelled_work,
                "ct_progress_status": ct_progress,
                "aseptic_progress_status": aseptic_progress,
                "epma_progress_status": epma_progress,
                # Kept for backwards compatibility: no elapsed time is credited
                # against CT planning allowance under the conservative rule.
                "ct_elapsed_bdays": 0,
                "ct_elapsed_since_start_bdays": int(ct_elapsed_context),
                "ct_remaining_bdays": int(ct_remaining),
                "aseptic_remaining_bdays": int(aseptic_remaining),
                "epma_remaining_bdays": int(epma_remaining),
                "aseptic_completed_by_milestone": bool(aseptic_completed),
                "epma_completed_by_milestone": bool(epma_completed),
                "progress_seen_by_planning_date": progress_seen,
                # Auditable milestone values actually used at this planning date.
                "observed_pharmacy_setup_start": setup_start,
                "observed_aseptic_notify": aseptic_notify,
                "observed_aseptic_bmr_authorised": aseptic_bmr,
                "observed_iqemo_setup": iqemo_setup,
                "observed_iqemo_authorised": iqemo_authorised,
                "observed_greenlight": greenlight,
            }
        )
        rows.append(output)

    return pd.DataFrame(rows)


# ============================================================
# RESOURCE AVAILABILITY VALIDATION
# ============================================================

def normalise_capacity_exceptions(rows, assumptions=None):
    assumptions = assumptions or DEFAULT_ASSUMPTIONS
    base_caps = {
        "CT": int(assumptions.get("cap_ct", 3)),
        "Aseptic": int(assumptions.get("cap_aseptic", 2)),
        "ePMA": int(assumptions.get("cap_epma", 1)),
    }
    cleaned = []
    for row in rows or []:
        if isinstance(row, pd.Series):
            row = row.to_dict()
        team = str(row.get("workstream", row.get("Workstream", ""))).strip()
        if team not in base_caps:
            continue
        start = _date_or_none(row.get("start_date", row.get("From")))
        end = _date_or_none(row.get("end_date", row.get("To")))
        if start is None or end is None:
            continue
        if end < start:
            raise ValueError(f"Resource availability end date is before start date for {team}.")
        cap = _number_or_none(row.get("available_capacity", row.get("Available capacity")))
        if cap is None:
            continue
        cap = int(round(cap))
        if cap < 0:
            raise ValueError("Available capacity cannot be negative.")
        if cap > base_caps[team]:
            raise ValueError(
                f"{team} temporary available capacity ({cap}) cannot exceed its base capacity ({base_caps[team]}) in the absence/reduced-capacity table."
            )
        cleaned.append(
            {
                "workstream": team,
                "start_date": start,
                "end_date": end,
                "available_capacity": cap,
                "note": _text_or_blank(row.get("note", row.get("Reason"))),
            }
        )
    return cleaned


# ============================================================
# SESSION STATE + PERSISTENCE
# ============================================================

def _jsonable(value):
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(v) for v in value]
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            pass
    return value


def _df_to_records(df):
    if df is None:
        return None
    out = df.copy()
    for col in out.columns:
        out[col] = out[col].map(_jsonable)
    return out.to_dict(orient="records")


def _restore_dataframe(records, date_columns=None):
    if records is None:
        return None
    df = pd.DataFrame(records)
    for col in date_columns or []:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce").dt.date
    return df


def serialise_working_state():
    return {
        "app_logic_version": APP_LOGIC_VERSION,
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "uploaded_tracker_name": st.session_state.get("uploaded_tracker_name"),
        "uploaded_source_type": st.session_state.get("uploaded_source_type"),
        "source_mapping": _jsonable(st.session_state.get("source_mapping", {})),
        "source_layout": _jsonable(st.session_state.get("source_layout", {})),
        "trial_table": _df_to_records(st.session_state.get("trial_table")),
        "assumptions": _jsonable(copy.deepcopy(st.session_state.get("assumptions", DEFAULT_ASSUMPTIONS))),
        "schedule_result": _df_to_records(st.session_state.get("schedule_result")),
        "schedule_info": _jsonable(st.session_state.get("schedule_info")),
        "schedule_dirty": bool(st.session_state.get("schedule_dirty", False)),
    }


def _apply_state_payload(data, bump_versions=True):
    table = _restore_dataframe(
        data.get("trial_table"),
        ["hra_date", "site_date", "site_confirmed_date", "pharmacy_notified_date", "pharmacy_setup_start_date"],
    )
    if table is not None:
        table = normalise_edited_trials(None, table, data.get("assumptions") or DEFAULT_ASSUMPTIONS)
    st.session_state.trial_table = table

    assumptions = data.get("assumptions") or copy.deepcopy(DEFAULT_ASSUMPTIONS)
    merged = copy.deepcopy(DEFAULT_ASSUMPTIONS)
    merged.update(assumptions)
    # One-time migration from the previous operational default.  Keep any custom
    # time budget, but upgrade the old 15-second default to the reproducible
    # 60-second default when loading a state created by an older logic version.
    if data.get("app_logic_version") != APP_LOGIC_VERSION:
        try:
            if int(merged.get("solver_time_limit_seconds", 15)) == 15:
                merged["solver_time_limit_seconds"] = 60
        except Exception:
            merged["solver_time_limit_seconds"] = 60
    if merged.get("planning_date"):
        merged["planning_date"] = pd.to_datetime(merged["planning_date"], errors="coerce").date()
    for team, tiers in merged.get("duration_tiers", {}).items():
        for tier, rng in list(tiers.items()):
            merged["duration_tiers"][team][tier] = tuple(rng)
    merged["capacity_exceptions"] = normalise_capacity_exceptions(
        merged.get("capacity_exceptions", []), merged
    )
    st.session_state.assumptions = merged

    # A saved schedule is valid only for the exact scheduling-logic version that
    # created it.  Trial data / assumptions can still be restored across code
    # updates, but derived optimiser outputs must be regenerated.
    saved_logic_version = data.get("app_logic_version")
    schedule_records = data.get("schedule_result") if saved_logic_version == APP_LOGIC_VERSION else None
    st.session_state.schedule_result = pd.DataFrame(schedule_records) if schedule_records is not None else None
    if st.session_state.schedule_result is not None:
        for col in [
            "release_date", "due_date", "pharmacy_notified_date", "scheduling_ready_date",
            "operational_release_date", "clock_start", "CT_start_rec", "CT_end_rec", "CT_capacity_end_rec",
            "Aseptic_start_rec", "Aseptic_end_rec", "ePMA_start_rec", "ePMA_end_rec",
            "completion_rec", "due_date_rec",
        ]:
            if col in st.session_state.schedule_result.columns:
                st.session_state.schedule_result[col] = pd.to_datetime(
                    st.session_state.schedule_result[col], errors="coerce"
                )

    st.session_state.schedule_info = data.get("schedule_info") if saved_logic_version == APP_LOGIC_VERSION else None
    st.session_state.schedule_dirty = (
        bool(data.get("schedule_dirty", False))
        if saved_logic_version == APP_LOGIC_VERSION
        else bool(st.session_state.trial_table is not None)
    )
    st.session_state.uploaded_tracker_name = data.get("uploaded_tracker_name")
    st.session_state.uploaded_source_type = data.get("uploaded_source_type")
    st.session_state.source_mapping = data.get("source_mapping") or {}
    st.session_state.source_layout = data.get("source_layout") or {}

    # Migrate notes/decisions/ranks from older saved-state format into the table.
    if st.session_state.trial_table is not None:
        t = st.session_state.trial_table.copy()
        old_comments = data.get("comments") or {}
        old_decisions = data.get("final_decisions") or {}
        old_ranks = data.get("manual_ranks") or {}
        for idx, row in t.iterrows():
            tid = str(row.get("trial_id", ""))
            if not _has_value(row.get("pharmacy_note")) and _has_value(old_comments.get(tid)):
                t.at[idx, "pharmacy_note"] = str(old_comments.get(tid))
            if not _has_value(row.get("final_decision")) and _has_value(old_decisions.get(tid)):
                t.at[idx, "final_decision"] = str(old_decisions.get(tid))
            if _number_or_none(row.get("manual_rank")) is None and _number_or_none(old_ranks.get(tid)) is not None:
                t.at[idx, "manual_rank"] = int(round(float(old_ranks.get(tid))))
        st.session_state.trial_table = normalise_edited_trials(None, t, st.session_state.assumptions)

    if bump_versions:
        st.session_state.trial_editor_version += 1
        st.session_state.priority_editor_version += 1
        st.session_state.schedule_editor_version += 1


def initialise_state():
    defaults = {
        "raw_df": None,
        "trial_table": None,
        "uploaded_tracker_name": None,
        "uploaded_source_type": None,
        "source_mapping": {},
        "source_layout": {},
        "assumptions": {**copy.deepcopy(DEFAULT_ASSUMPTIONS), "planning_date": current_planning_date()},
        "schedule_result": None,
        "schedule_info": None,
        "schedule_dirty": False,
        "trial_editor_version": 0,
        "priority_editor_version": 0,
        "schedule_editor_version": 0,
        "_autosave_checked": False,
        "_last_autosave_at": None,
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value

    # Invalidate any in-memory schedule produced by an older code version.
    # This is especially important in Streamlit because session_state and local
    # autosave can otherwise preserve old derived values after .py files change.
    previous_logic_version = st.session_state.get("_app_logic_version")
    if previous_logic_version != APP_LOGIC_VERSION:
        st.session_state.schedule_result = None
        st.session_state.schedule_info = None
        st.session_state.schedule_dirty = bool(st.session_state.trial_table is not None)
        # Upgrade only the old default; preserve any deliberately customised value.
        try:
            if int(st.session_state.assumptions.get("solver_time_limit_seconds", 15)) == 15:
                st.session_state.assumptions["solver_time_limit_seconds"] = 60
        except Exception:
            st.session_state.assumptions["solver_time_limit_seconds"] = 60
        # v8 operational default: specialist branches become available two working
        # days after CT starts. Earlier app versions used 0 as the default, so migrate
        # that legacy default once; staff can still manually set 0 afterwards.
        try:
            if int(st.session_state.assumptions.get("specialist_handoff_bdays", 0)) == 0:
                st.session_state.assumptions["specialist_handoff_bdays"] = 2
        except Exception:
            st.session_state.assumptions["specialist_handoff_bdays"] = 2
        st.session_state["_app_logic_version"] = APP_LOGIC_VERSION

    # Resume the last locally saved editable workspace in a new browser session.
    # Operational safety rule: an automatic resume restores the trial data, notes,
    # statuses and resource assumptions, but the planning date resets to TODAY.
    # A historical planning date is restored only when a user explicitly loads a
    # named checkpoint / JSON backup. This prevents an old Chapter-6-style
    # snapshot date from silently becoming the live planning date on reopening.
    if not st.session_state.get("_autosave_checked", False):
        st.session_state._autosave_checked = True
        if not PUBLIC_DEMO and st.session_state.trial_table is None and AUTO_SAVE_FILE.exists():
            try:
                data = json.loads(AUTO_SAVE_FILE.read_text(encoding="utf-8"))
                _apply_state_payload(data, bump_versions=False)
                if data.get("app_logic_version") != APP_LOGIC_VERSION:
                    try:
                        if int(st.session_state.assumptions.get("specialist_handoff_bdays", 0)) == 0:
                            st.session_state.assumptions["specialist_handoff_bdays"] = 2
                            st.session_state.schedule_result = None
                            st.session_state.schedule_info = None
                            st.session_state.schedule_dirty = True
                    except Exception:
                        st.session_state.assumptions["specialist_handoff_bdays"] = 2
                        st.session_state.schedule_result = None
                        st.session_state.schedule_info = None
                        st.session_state.schedule_dirty = True
                restored_date = st.session_state.assumptions.get("planning_date")
                today = current_planning_date()
                if restored_date != today:
                    st.session_state.assumptions["planning_date"] = today
                    # A schedule calculated for a previous planning date must not
                    # be presented as current without recalculation.
                    if st.session_state.schedule_result is not None:
                        st.session_state.schedule_dirty = True
                st.session_state._last_autosave_at = data.get("saved_at")
            except Exception:
                pass


def autosave_working_state():
    if PUBLIC_DEMO:
        return None
    if "trial_table" not in st.session_state or st.session_state.get("trial_table") is None:
        return None
    SAVED_STATES_DIR.mkdir(parents=True, exist_ok=True)
    payload = serialise_working_state()
    payload["name"] = "Automatic workspace"
    AUTO_SAVE_FILE.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    st.session_state._last_autosave_at = payload.get("saved_at")
    return AUTO_SAVE_FILE


def clear_autosave():
    if AUTO_SAVE_FILE.exists():
        AUTO_SAVE_FILE.unlink()
    st.session_state._last_autosave_at = None


def mark_schedule_dirty(auto_save=True):
    st.session_state["schedule_dirty"] = True
    if auto_save:
        try:
            autosave_working_state()
        except Exception:
            pass


def mark_schedule_clean(auto_save=True):
    st.session_state["schedule_dirty"] = False
    if auto_save:
        try:
            autosave_working_state()
        except Exception:
            pass


def save_working_state(name):
    clean_name = str(name or "").strip()
    if not clean_name:
        raise ValueError("Enter a name for the saved state.")
    SAVED_STATES_DIR.mkdir(parents=True, exist_ok=True)
    state_id = uuid.uuid4().hex[:10]
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", clean_name).strip("_") or "saved_state"
    payload = serialise_working_state()
    payload["name"] = clean_name
    payload["state_id"] = state_id
    path = SAVED_STATES_DIR / f"{safe}_{state_id}.json"
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def list_saved_working_states():
    if not SAVED_STATES_DIR.exists():
        return []
    states = []
    for path in SAVED_STATES_DIR.glob("*.json"):
        if path.name == AUTO_SAVE_FILE.name:
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            states.append({"file": path.name, "name": data.get("name", path.stem), "saved_at": data.get("saved_at", "")})
        except Exception:
            continue
    return sorted(states, key=lambda x: x.get("saved_at", ""), reverse=True)


def load_working_state(filename):
    path = SAVED_STATES_DIR / Path(filename).name
    if not path.exists():
        raise FileNotFoundError("The selected saved state no longer exists.")
    data = json.loads(path.read_text(encoding="utf-8"))
    _apply_state_payload(data, bump_versions=True)
    autosave_working_state()


def load_working_state_from_json_bytes(data_bytes: bytes):
    data = json.loads(data_bytes.decode("utf-8"))
    _apply_state_payload(data, bump_versions=True)
    autosave_working_state()


def delete_saved_working_state(filename):
    path = SAVED_STATES_DIR / Path(filename).name
    if path.exists():
        path.unlink()


def current_workspace_json_bytes():
    return json.dumps(serialise_working_state(), indent=2, ensure_ascii=False).encode("utf-8")


def current_trial_export_dataframe():
    """Return a reloadable, human-readable snapshot of the editable trial table."""
    table = st.session_state.get("trial_table")
    if table is None:
        return pd.DataFrame()
    t = table.copy().reset_index(drop=True)
    out = pd.DataFrame({
        "Trial": t.get("trial_id", pd.Series(dtype=object)),
        "Trial status": t.get("trial_status", pd.Series(dtype=object)),
        "HRA approval": t.get("hra_date", pd.Series(dtype=object)),
        "Site selected": t.get("site_date", pd.Series(dtype=object)),
        "Pharmacy notified": t.get("pharmacy_notified_date", pd.Series(dtype=object)),
        "Pharmacy set-up start": t.get("pharmacy_setup_start_date", pd.Series(dtype=object)),
        "Pharmacy CT notify Aseptic Services": t.get("aseptic_notify_date", pd.Series(dtype=object)),
        "Date Aseptic BMR Authorised": t.get("aseptic_bmr_authorised_date", pd.Series(dtype=object)),
        "iQemo Template set up": t.get("iqemo_setup_date", pd.Series(dtype=object)),
        "iQemo template authorised": t.get("iqemo_authorised_date", pd.Series(dtype=object)),
        "Pharmacy Greenlight Approval": t.get("greenlight_date", pd.Series(dtype=object)),
        "Aseptic": t.get("aseptic_required", pd.Series(dtype=object)),
        "Prescribing system": t.get("prescribing_system", pd.Series(dtype=object)),
        "CT Complexity": t.get("ct_complexity", pd.Series(dtype=object)),
        "ePMA Complexity": t.get("epma_complexity", pd.Series(dtype=object)),
        "Stage 2 priority score": t.get("priority_score", pd.Series(dtype=object)),
        "Manual rank": t.get("manual_rank", pd.Series(dtype=object)),
        "Final decision": t.get("final_decision", pd.Series(dtype=object)),
        "Pharmacy Note": t.get("pharmacy_note", pd.Series(dtype=object)),
    })
    # Useful context fields are retained when present, without making them required.
    optional = [
        ("Date site confirmed (EDGE)", "site_confirmed_date"),
        ("Trial phase", "trial_phase"),
        ("Q4b pharmacy & aseptics burden", "q4b_burden"),
    ]
    for label, col in optional:
        if col in t.columns:
            out[label] = t[col]
    return out


def _style_excel_sheet(ws, date_headers=None, preferred_widths=None):
    """Apply readable widths and DD/MM/YYYY formats to an openpyxl worksheet."""
    from openpyxl.utils import get_column_letter

    date_headers = set(date_headers or [])
    preferred_widths = preferred_widths or {}
    headers = {cell.value: cell.column for cell in ws[1] if cell.value is not None}

    for header, col_idx in headers.items():
        letter = get_column_letter(col_idx)
        if header in date_headers:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value not in (None, ""):
                    cell.number_format = "DD/MM/YYYY"

        if header in preferred_widths:
            width = preferred_widths[header]
        else:
            max_len = len(str(header))
            for row in range(2, min(ws.max_row, 250) + 1):
                value = ws.cell(row=row, column=col_idx).value
                if value is None:
                    continue
                if header in date_headers:
                    text = "DD/MM/YYYY"
                else:
                    text = str(value)
                max_len = max(max_len, min(len(text), 60))
            width = min(max(max_len + 2, 12), 42)
        ws.column_dimensions[letter].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def dataframe_to_excel_bytes(df: pd.DataFrame, sheet_name="Data", date_columns=None, preferred_widths=None):
    """Export a dataframe as a readable Excel sheet using UK date formatting."""
    output = io.BytesIO()
    date_columns = [c for c in (date_columns or []) if c in df.columns]
    clean = dates_as_date_objects(df, date_columns)
    with pd.ExcelWriter(
        output,
        engine="openpyxl",
        date_format="DD/MM/YYYY",
        datetime_format="DD/MM/YYYY",
    ) as writer:
        clean.to_excel(writer, index=False, sheet_name=sheet_name)
        _style_excel_sheet(
            writer.book[sheet_name],
            date_headers=date_columns,
            preferred_widths=preferred_widths,
        )
    return output.getvalue()


def current_trial_workbook_bytes():
    """Create a reloadable, readable Excel export of current trial edits/settings."""
    output = io.BytesIO()
    trial_df = current_trial_export_dataframe()
    trial_date_cols = [
        "HRA approval", "Site selected", "Pharmacy notified",
        "Pharmacy set-up start", "Pharmacy CT notify Aseptic Services",
        "Date Aseptic BMR Authorised", "iQemo Template set up",
        "iQemo template authorised", "Pharmacy Greenlight Approval",
        "Date site confirmed (EDGE)",
    ]
    trial_df = dates_as_date_objects(trial_df, trial_date_cols)

    assumptions = st.session_state.get("assumptions") or DEFAULT_ASSUMPTIONS
    planning_date = _date_or_none(assumptions.get("planning_date"))
    assumption_rows = [
        {"Setting": "Planning date", "Value": planning_date},
        {"Setting": "55-day target", "Value": assumptions.get("target_days")},
        {"Setting": "CT base capacity", "Value": assumptions.get("cap_ct")},
        {"Setting": "CT capacity accounting", "Value": assumptions.get("ct_capacity_mode", "continuous")},
        {"Setting": "Aseptic base capacity", "Value": assumptions.get("cap_aseptic")},
        {"Setting": "ePMA base capacity", "Value": assumptions.get("cap_epma")},
        {"Setting": "Aseptic default working days", "Value": assumptions.get("aseptic_default_bdays")},
    ]
    assumptions_df = pd.DataFrame(assumption_rows)

    exceptions = normalise_capacity_exceptions(
        assumptions.get("capacity_exceptions", []), assumptions
    )
    resources_df = pd.DataFrame(exceptions) if exceptions else pd.DataFrame(
        columns=["workstream", "start_date", "end_date", "available_capacity", "note"]
    )
    resources_df = resources_df.rename(columns={
        "workstream": "Workstream",
        "start_date": "From",
        "end_date": "To",
        "available_capacity": "Available capacity",
        "note": "Reason",
    })
    resources_df = dates_as_date_objects(resources_df, ["From", "To"])

    with pd.ExcelWriter(
        output, engine="openpyxl", date_format="DD/MM/YYYY", datetime_format="DD/MM/YYYY"
    ) as writer:
        trial_df.to_excel(writer, index=False, sheet_name="Trials")
        assumptions_df.to_excel(writer, index=False, sheet_name="Assumptions")
        resources_df.to_excel(writer, index=False, sheet_name="Resource availability")

        _style_excel_sheet(
            writer.book["Trials"],
            date_headers=trial_date_cols,
            preferred_widths={
                "Trial": 28, "Trial status": 16, "Prescribing system": 22,
                "CT Complexity": 16, "ePMA Complexity": 18,
                "Final decision": 28, "Pharmacy Note": 55,
            },
        )
        _style_excel_sheet(
            writer.book["Assumptions"],
            date_headers=[],
            preferred_widths={"Setting": 32, "Value": 20},
        )
        # Only From/To are date columns on the resource sheet.
        _style_excel_sheet(
            writer.book["Resource availability"],
            date_headers=["From", "To"],
            preferred_widths={"Workstream": 18, "From": 14, "To": 14, "Available capacity": 20, "Reason": 40},
        )

        # Assumptions has mixed date/numeric values in one column; format only the
        # planning-date row as a UK date instead of the whole column.
        ws_a = writer.book["Assumptions"]
        if planning_date is not None and ws_a.max_row >= 2:
            ws_a.cell(row=2, column=2).number_format = "DD/MM/YYYY"

    return output.getvalue()


def render_last_recalculated():
    info = st.session_state.get("schedule_info") or {}
    calculated_at = info.get("calculated_at")
    if calculated_at:
        ts = pd.Timestamp(calculated_at)
        text = f"Last recalculated: {ts.strftime('%d %b %Y, %H:%M')}"
    else:
        text = "Last recalculated: Not yet calculated"
    if st.session_state.get("schedule_dirty") and calculated_at:
        text += " · inputs changed since last calculation"
    st.caption(text)


def load_sample_workspace():
    """Load the synthetic demo portfolio into the current session."""
    from sample_data import build_sample_dataframe

    df = build_sample_dataframe(current_planning_date())
    resolved = resolve_tracker_columns(df)
    mapping = {k: v for k, v in resolved.items() if v}
    df.attrs["column_mapping"] = mapping
    table = build_operational_trial_table(df, st.session_state.assumptions, mapping=mapping)
    st.session_state.raw_df = df
    st.session_state.trial_table = table
    st.session_state.uploaded_tracker_name = "Synthetic sample portfolio"
    st.session_state.uploaded_source_type = "synthetic sample"
    st.session_state.source_mapping = mapping
    st.session_state.source_layout = {"sheet_name": "Trials", "header_row_excel": 1}
    st.session_state.schedule_result = None
    st.session_state.schedule_info = None
    st.session_state.trial_editor_version += 1
    mark_schedule_dirty(auto_save=True)
    return table


def render_sample_data_button(key="load_sample"):
    if st.button("Load synthetic sample data", type="primary", key=key,
                 help="Loads 20 fictitious trials so you can try every page without uploading a file."):
        try:
            load_sample_workspace()
            st.rerun()
        except Exception as exc:
            st.error(f"Could not load the sample data: {exc}")


def render_saved_work_sidebar():
    if PUBLIC_DEMO:
        with st.sidebar.expander("💾 Workspace", expanded=False):
            st.caption("Demo mode: nothing is stored on the server. Use **Review** to download a backup of your workspace.")
        return
    with st.sidebar.expander("💾 Workspace", expanded=False):
        auto = st.session_state.get("_last_autosave_at")
        if auto:
            st.caption("Auto-saved locally: " + str(auto).replace("T", " ")[:16])
        else:
            st.caption("Changes are auto-saved locally after edits.")

        if st.button(
            "Save changes now",
            use_container_width=True,
            disabled=st.session_state.get("trial_table") is None,
            key="sidebar_save_changes_now",
        ):
            try:
                autosave_working_state()
                st.success("Current workspace saved locally.")
            except Exception as exc:
                st.error(str(exc))

        save_name = st.text_input("Named checkpoint", placeholder="e.g. Monday pharmacy review", key="saved_work_name")
        if st.button("Save checkpoint", use_container_width=True, disabled=st.session_state.get("trial_table") is None):
            try:
                save_working_state(save_name)
                st.success("Checkpoint saved.")
            except Exception as exc:
                st.error(str(exc))

        states = list_saved_working_states()
        if states:
            label_map = {f"{x['name']} · {x['saved_at'].replace('T', ' ')[:16]}": x["file"] for x in states}
            selected = st.selectbox("Saved checkpoints", list(label_map), key="saved_work_selector")
            c1, c2 = st.columns(2)
            if c1.button("Load", use_container_width=True):
                load_working_state(label_map[selected])
                st.rerun()
            if c2.button("Delete", use_container_width=True):
                delete_saved_working_state(label_map[selected])
                st.rerun()
